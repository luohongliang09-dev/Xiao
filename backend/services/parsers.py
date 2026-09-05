import io

from docx import Document


def parse_txt_md(data: bytes) -> str:
    """txt / markdown：按编码顺序尝试解码。"""
    for enc in ("utf-8", "utf-8-sig", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def parse_docx(data: bytes) -> str:
    """Word .docx：抽取段落 + 表格。"""
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    import pymupdf as fitz

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        pages = []
        for page in doc:
            t = page.get_text().strip()
            if t:
                pages.append(t)
        return "\n\n".join(pages)
    finally:
        doc.close()


def _pdf_ocr(data: bytes) -> str:
    """扫描版 PDF：渲染成图片后用 OCR 识别。"""
    import numpy as np
    import pymupdf as fitz
    from rapidocr_onnxruntime import RapidOCR

    ocr = RapidOCR()
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        texts = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # 2x 缩放 ≈ 144dpi
            img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 4:
                img = img[:, :, :3].copy()
            result, _ = ocr(img)
            if result:
                texts.append("\n".join(line[1] for line in result))
        return "\n\n".join(texts)
    finally:
        doc.close()


def parse_pdf(data: bytes) -> str:
    text = _pdf_text(data)
    if text.strip():
        return text
    # 文字版抽不到内容 → 判定为扫描版，走 OCR
    try:
        return _pdf_ocr(data)
    except Exception:
        return ""


def parse_document(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    return parse_txt_md(data)
